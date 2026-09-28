"""Builds assets/templates/control-testing-workpaper.xlsx (test plan, sample results, conclusion, sample size guide).
Run from the repo root:  python3 _scripts/templates/build_workpaper.py   (then recalculate with LibreOffice; see README)
"""
from pathlib import Path
import sys, datetime
from math import comb
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import (f, header_row, style_cell, list_validation, title_block, print_setup,
                         INPUT_FILL, CALC_FILL, EXAMPLE_FILL, DISCLAIMER, NAVY)

ROOT = Path(__file__).resolve().parents[2]
V = "version 1.0 (September 2026)"
SITE = "https://hayesskvaril-ctrl.github.io"
D = datetime.date
ROWS = 60

wb = Workbook()

# ---------------------------------------------------------------- Read me
ws = wb.active
ws.title = "Read me"
title_block(ws, "Control testing workpaper", f"RiskLens Australia · free template · {V}")
readme = [
    ("What this is", f"A workpaper for testing one control: the test plan, design assessment, sample results and conclusion, with reviewer sign-off. Background: {SITE}/risk-management/control-design-and-testing.html and {SITE}/risk-management/control-testing-sampling.html"),
    ("How to use it", "1. Test plan: describe the control, the risk it addresses, the population and the attributes you will test (up to four). Record the design assessment.\n"
                      "2. Choose a sample size using your methodology (the 'Sample size guide' tab shows statistical sizes for attribute sampling).\n"
                      "3. Sample results: record each item tested and whether each attribute was met (Yes, No or N/A). Exceptions are flagged automatically.\n"
                      "4. Conclusion: review the exception count, exception rate and upper deviation limit, then record the operating effectiveness conclusion, findings and sign-offs."),
    ("Worked example", "The yellow cells hold a fictional example: a four-eyes approval control over super benefit payments, with 25 items tested and one exception. Replace it with your own test."),
    ("Upper deviation limit", "The Conclusion tab calculates how high the true exception rate could be, at the confidence level you choose, given what the sample found. It uses an exact binomial (Clopper-Pearson) upper bound and assumes a random sample from a large population. It's a statistical aid, not a substitute for judgement: one exception that reveals a design flaw or fraud matters more than the rate suggests."),
    ("Colour key", "Pale yellow cells: fill these in (many have drop-down lists).\nGrey-blue cells: calculated automatically; don't type over them.\nGreen rows: worked examples. Replace or delete them."),
    ("Important", DISCLAIMER),
]
for i, (k, v) in enumerate(readme):
    r = 4 + i
    ws.cell(row=r, column=1, value=k).font = f(11, True, NAVY)
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
    c = ws.cell(row=r, column=2, value=v)
    c.font = f(10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = max(30, 15 * (v.count("\n") + 1 + len(v) // 95))
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 100

# ---------------------------------------------------------------- Lists
ls = wb.create_sheet("Lists")
title_block(ls, "Drop-down lists", "Edit these values to suit your organisation. The drop-downs read from this tab.")
LISTS = {
    "Frequency": ["Multiple times a day", "Daily", "Weekly", "Monthly", "Quarterly", "Annually", "Event-driven"],
    "Control type": ["Preventive", "Detective", "Corrective"],
    "Manual or automated": ["Manual", "IT-dependent manual", "Automated"],
    "Test method": ["Inquiry only", "Inquiry and observation", "Inspection of evidence", "Reperformance", "Combination"],
    "Design": ["Effective", "Partially effective", "Ineffective"],
    "Result": ["Yes", "No", "N/A"],
    "Conclusion": ["Effective", "Partially effective", "Ineffective", "Unable to conclude"],
    "Confidence": [0.8, 0.9, 0.95, 0.99],
    "Yes / No": ["Yes", "No"],
}
names = list(LISTS)
for j, n in enumerate(names):
    col = j + 1
    style_cell(ls.cell(row=4, column=col, value=n), bold=True)
    for i, v in enumerate(LISTS[n]):
        style_cell(ls.cell(row=5 + i, column=col, value=v), INPUT_FILL, fmt="0%" if n == "Confidence" else None)
    ls.column_dimensions[chr(64 + col)].width = 24

def ref(n):
    col = chr(65 + names.index(n))
    return f"=Lists!${col}$5:${col}${4 + len(LISTS[n])}"

# ---------------------------------------------------------------- Test plan (form layout)
tp = wb.create_sheet("Test plan", 1)
title_block(tp, "Test plan and design assessment", "One workpaper per control. Fill in the yellow cells.")
tp.column_dimensions["A"].width = 34
tp.column_dimensions["B"].width = 80
PLAN = [  # (label, example value, list name or None, number format)
    ("SECTION", "About the control"),
    ("Workpaper reference", "CT-2026-07", None, None),
    ("Control ID", "C-PAY-03", None, None),
    ("Control description", "A second officer reviews and approves every benefit payment over $10,000 in the payment system before release, checking the payment against the member's claim form, identity verification and bank details.", None, None),
    ("Control objective", "Benefit payments over $10,000 are valid, accurate and paid to the right member.", None, None),
    ("Risk addressed", "Payment made to the wrong person or in the wrong amount (error or fraud).", None, None),
    ("Control owner", "Payments Team Leader", None, None),
    ("Frequency", "Multiple times a day", "Frequency", None),
    ("Control type", "Preventive", "Control type", None),
    ("Manual or automated", "IT-dependent manual", "Manual or automated", None),
    ("Key control?", "Yes", "Yes / No", None),
    ("SECTION", "Design assessment"),
    ("Does the design address the risk? (reasoning)", "Yes. The control operates before payment release, covers all payments over the threshold, and the system prevents the same user from creating and approving a payment.", None, None),
    ("Design gaps identified", "Payments just under $10,000 are not reviewed; split-payment monitoring relies on a separate detective report.", None, None),
    ("Design conclusion", "Effective", "Design", None),
    ("SECTION", "Test approach"),
    ("Test period from", D(2026, 1, 1), None, "DD/MM/YYYY"),
    ("Test period to", D(2026, 6, 30), None, "DD/MM/YYYY"),
    ("Population description", "All benefit payments over $10,000 released between 1 January and 30 June 2026 (system report PAY-R12).", None, None),
    ("Population size", 4830, None, "#,##0"),
    ("How completeness of the population was checked", "Reconciled report total to the general ledger benefit payments account for the period.", None, None),
    ("Test method", "Reperformance", "Test method", None),
    ("Sampling method", "Random selection using a random number generator over the report row numbers.", None, None),
    ("Planned sample size", 25, None, None),
    ("Basis for sample size", "Methodology table for controls operating multiple times a day. See the Sample size guide tab for the confidence this implies.", None, None),
    ("Attribute A", "Approver is a different person from the payment creator.", None, None),
    ("Attribute B", "Payment amount agrees to the claim form and benefit calculation.", None, None),
    ("Attribute C", "Bank details agree to verified member records.", None, None),
    ("Attribute D", "Approval evidenced in the system before payment release.", None, None),
    ("Tester", "Operational Risk Analyst", None, None),
    ("Date testing completed", D(2026, 7, 24), None, "DD/MM/YYYY"),
]
cells = {}
r = 4
for item in PLAN:
    if item[0] == "SECTION":
        c = tp.cell(row=r, column=1, value=item[1]); c.font = f(12, True, NAVY)
        r += 1
        continue
    label, val, lst, fmt = item
    style_cell(tp.cell(row=r, column=1, value=label), bold=True)
    c = tp.cell(row=r, column=2, value=val)
    style_cell(c, INPUT_FILL, fmt=fmt)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    if lst:
        list_validation(tp, ref(lst), f"B{r}")
    if isinstance(val, str) and len(val) > 70:
        tp.row_dimensions[r].height = 15 * (1 + len(val) // 75)
    cells[label] = f"B{r}"
    r += 1

# ---------------------------------------------------------------- Sample results
sr = wb.create_sheet("Sample results", 2)
title_block(sr, "Sample results", "One row per item tested. Record Yes (met), No (not met) or N/A for each attribute.")
heads = ["#", "Item reference", "Item date", "Attribute A", "Attribute B", "Attribute C", "Attribute D", "Exception?", "Exception description", "Evidence reference", "Tester notes"]
widths = [5, 18, 12, 14, 14, 14, 14, 11, 40, 18, 30]
header_row(sr, 4, heads, widths, height=30)
# Attribute headers show the attribute text from the test plan as a second header row
style_cell(sr.cell(row=5, column=1, value="Tests:"), CALC_FILL, bold=True)
for k, att in zip("DEFG", ["Attribute A", "Attribute B", "Attribute C", "Attribute D"]):
    c = sr[f"{k}5"]; c.value = f"='Test plan'!{cells[att]}"
    style_cell(c, CALC_FILL); c.font = f(8, False, "4A5568", italic=True)
sr.row_dimensions[5].height = 60
for k in "BCHIJK":
    style_cell(sr[f"{k}5"], CALC_FILL)
FIRST, LAST = 6, 5 + ROWS
ex_rows = []
for i in range(25):
    ok = ["Yes", "Yes", "Yes", "Yes"]
    desc = note = None
    if i == 13:
        ok = ["Yes", "Yes", "Yes", "No"]
        desc = "Approval recorded 40 minutes after payment release. Payment was valid, but the control did not operate before release."
        note = "Confirmed with Payments Team Leader: a system override was used during an outage. Override log reviewed."
    ex_rows.append([i + 1, f"PAY-{260000 + 173 * (i + 1) % 9973:06d}", D(2026, 1 + i // 5, 3 + (i * 5) % 25), *ok, desc, f"Screenshot folder CT-2026-07/{i + 1:02d}", note])
for rr in range(FIRST, LAST + 1):
    ex = ex_rows[rr - FIRST] if rr - FIRST < len(ex_rows) else None
    for ci in range(1, 12):
        cell = sr.cell(row=rr, column=ci)
        if ci == 8:
            cell.value = f'=IF(COUNTA(D{rr}:G{rr})=0,"",IF(COUNTIF(D{rr}:G{rr},"No")>0,"Yes","No"))'
            style_cell(cell, CALC_FILL, center=True)
            continue
        if ci == 1:
            cell.value = rr - FIRST + 1
            style_cell(cell, CALC_FILL, center=True)
            continue
        if ex is not None:
            v = ex[ci - 1] if ci < 8 else ex[ci - 2]
            cell.value = v
        style_cell(cell, EXAMPLE_FILL if ex is not None else INPUT_FILL, fmt="DD/MM/YYYY" if ci == 3 else None, center=4 <= ci <= 7)
list_validation(sr, ref("Result"), f"D{FIRST}:G{LAST}")
for rng in (f"D{FIRST}:G{LAST}", f"H{FIRST}:H{LAST}"):
    first = rng.split(":")[0]
    sr.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="No"' if rng.startswith("D") else f'{first}="Yes"'], fill=PatternFill("solid", fgColor="FCA5A5")))
sr.freeze_panes = f"B{FIRST}"
sr.auto_filter.ref = f"A4:K{LAST}"

# ---------------------------------------------------------------- Conclusion
cn = wb.create_sheet("Conclusion", 3)
title_block(cn, "Results and conclusion", "Calculated results update from the Sample results tab. Fill in the yellow cells.")
cn.column_dimensions["A"].width = 44
cn.column_dimensions["B"].width = 70
S = f"'Sample results'!"
rows = [
    ("SECTION", "Results"),
    ("Items tested", f"=COUNTIF({S}H{FIRST}:H{LAST},\"Yes\")+COUNTIF({S}H{FIRST}:H{LAST},\"No\")", "calc", "0"),
    ("Planned sample size", f"='Test plan'!{cells['Planned sample size']}", "calc", "0"),
    ("Exceptions (items with at least one 'No')", f"=COUNTIF({S}H{FIRST}:H{LAST},\"Yes\")", "calc", "0"),
    ("Exception rate in the sample", "=IF(B5=0,\"\",B7/B5)", "calc", "0.0%"),
    ("Exceptions by attribute (A / B / C / D)", f"=COUNTIF({S}D{FIRST}:D{LAST},\"No\")&\" / \"&COUNTIF({S}E{FIRST}:E{LAST},\"No\")&\" / \"&COUNTIF({S}F{FIRST}:F{LAST},\"No\")&\" / \"&COUNTIF({S}G{FIRST}:G{LAST},\"No\")", "calc", None),
    ("SECTION", "Statistical check (optional)"),
    ("Confidence level", 0.95, "Confidence", "0%"),
    ("Tolerable exception rate", 0.1, "in", "0%"),
    ("Upper deviation limit", "=IF(B5=0,\"\",IF(B7>=B5,1,BETAINV(B11,B7+1,B5-B7)))", "calc", "0.0%"),
    ("Statistical result", "=IF(B13=\"\",\"\",IF(B13<=B12,\"Supports reliance: upper limit is within the tolerable rate\",\"Does not support reliance: upper limit exceeds the tolerable rate\"))", "calc", None),
    ("Sample complete?", "=IF(B5=0,\"\",IF(B5>=B6,\"Yes\",\"No: \"&(B6-B5)&\" items still to test\"))", "calc", None),
    ("SECTION", "Conclusion"),
    ("Root cause of exceptions", "A system override used during a payment platform outage let a payment release before approval. Override use was not reviewed after the event.", "in", None),
    ("Is an exception isolated or systemic? (reasoning)", "Isolated to one outage day, but the override process is a design gap: it could recur in any outage.", "in", None),
    ("Operating effectiveness conclusion", "Partially effective", "Conclusion", None),
    ("Overall control rating", "Partially effective", "Conclusion", None),
    ("Findings and agreed actions", "1. Require retrospective approval and second-line review of every override (owner: Payments Team Leader, due 30/09/2026).\n2. Add override use to the monthly payments KRI report.", "in", None),
    ("Issue raised in issues register?", "Yes", "Yes / No", None),
    ("SECTION", "Sign-off"),
    ("Prepared by", "Operational Risk Analyst", "in", None),
    ("Date prepared", D(2026, 7, 24), "in", "DD/MM/YYYY"),
    ("Reviewed by", "Operational Risk Manager", "in", None),
    ("Date reviewed", D(2026, 7, 29), "in", "DD/MM/YYYY"),
    ("Reviewer comments", "Agree with conclusion. Confirm override review is added to the control description once implemented.", "in", None),
]
r = 4
for item in rows:
    if item[0] == "SECTION":
        cn.cell(row=r, column=1, value=item[1]).font = f(12, True, NAVY)
        r += 1
        continue
    label, val, kind, fmt = item
    style_cell(cn.cell(row=r, column=1, value=label), bold=True)
    c = cn.cell(row=r, column=2, value=val)
    style_cell(c, CALC_FILL if kind == "calc" else INPUT_FILL, fmt=fmt)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    if kind not in ("calc", "in"):
        list_validation(cn, ref(kind), f"B{r}")
    if isinstance(val, str) and not val.startswith("=") and len(val) > 70:
        cn.row_dimensions[r].height = 15 * (1 + len(val) // 70 + val.count("\n"))
    r += 1
# sanity: the formulas above refer to fixed rows; check the layout matches
assert cn["A5"].value == "Items tested" and cn["A7"].value.startswith("Exceptions (") and cn["A11"].value == "Confidence level" and cn["A13"].value == "Upper deviation limit"
cn.conditional_formatting.add("B14", FormulaRule(formula=['LEFT(B14,4)="Does"'], fill=PatternFill("solid", fgColor="FCA5A5")))
cn.conditional_formatting.add("B14", FormulaRule(formula=['LEFT(B14,4)="Supp"'], fill=PatternFill("solid", fgColor="BFE5CB")))

# ---------------------------------------------------------------- Sample size guide
sg = wb.create_sheet("Sample size guide", 4)
title_block(sg, "Sample size guide (attribute sampling)", "Minimum sample sizes for a large population, from the binomial distribution.")
def n_for(conf, p, k):
    n = k + 1
    while sum(comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1)) > 1 - conf:
        n += 1
    return n
header_row(sg, 4, ["Confidence", "Tolerable exception rate", "0 exceptions allowed", "1 exception allowed", "2 exceptions allowed"], [14, 16, 14, 14, 14])
r = 5
for conf in (0.8, 0.9, 0.95):
    for p in (0.05, 0.1, 0.15, 0.2):
        vals = [conf, p] + [n_for(conf, p, k) for k in range(3)]
        for ci, v in enumerate(vals, start=1):
            style_cell(sg.cell(row=r, column=ci, value=v), CALC_FILL, fmt="0%" if ci <= 2 else "0", center=True)
        r += 1
notes = [
    "How to read it: to be 95% confident the true exception rate is no more than 10%, test 29 items and find no exceptions, or 46 items and find no more than one.",
    "These sizes assume a random sample from a large population. For small populations (e.g. a quarterly control with four occurrences), many testers test all or most occurrences instead.",
    "Many organisations set sample sizes by control frequency in their testing methodology. Those tables are practical; this guide helps show the level of confidence they imply.",
    f"More explanation and an interactive calculator: {SITE}/risk-management/control-testing-sampling.html",
]
for i, t in enumerate(notes):
    c = sg.cell(row=r + 1 + i, column=1, value=t); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical="top")
    sg.merge_cells(start_row=r + 1 + i, start_column=1, end_row=r + 1 + i, end_column=5)
    sg.row_dimensions[r + 1 + i].height = 32

print_setup(wb)
out = ROOT / "assets" / "templates" / "control-testing-workpaper.xlsx"
wb.save(out)
print("wrote", out.relative_to(ROOT))
