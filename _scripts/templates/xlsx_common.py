"""Shared styling helpers for the RiskLens spreadsheet templates (openpyxl)."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY = "0F2942"
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")   # pale yellow: fill in
CALC_FILL = PatternFill("solid", fgColor="EEF2F7")    # grey-blue: calculated
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
EXAMPLE_FILL = PatternFill("solid", fgColor="E7F6EC")  # green tint: example row
THIN = Side(style="thin", color="C9D3E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RATING_COLOURS = {"Low": "BFE5CB", "Medium": "FDE68A", "High": "FDBA74", "Extreme": "FCA5A5"}

DISCLAIMER = ("RiskLens Australia provides general educational information only. It is not legal, financial or "
              "compliance advice and does not take into account your circumstances. Always refer to the official "
              "legislation and regulator guidance, and seek professional advice where appropriate.")


def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def header_row(ws, row, headers, widths=None, height=32):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f(10, True, "FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    ws.row_dimensions[row].height = height
    if widths:
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(c, fill=None, wrap=True, bold=False, fmt=None, center=False):
    c.font = f(10, bold)
    if fill is not None:
        c.fill = fill
    c.border = BORDER
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if center else None)
    if fmt:
        c.number_format = fmt


def list_validation(ws, formula, rng, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    if prompt:
        dv.promptTitle = "Choose from the list"
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.error = "Please choose a value from the drop-down list."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def rating_colours(ws, rng):
    """Colour cells whose text is Low/Medium/High/Extreme."""
    first = rng.split(":")[0]
    for name, colour in RATING_COLOURS.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{name}"'],
                                                       fill=PatternFill("solid", fgColor=colour)))


def title_block(ws, title, subtitle, width_cols=8):
    ws["A1"] = title
    ws["A1"].font = f(16, True, NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = f(10, False, "4A5568", italic=True)
    ws.row_dimensions[1].height = 26


def print_setup(wb):
    """Landscape, fit all columns to one page wide, for every sheet."""
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
