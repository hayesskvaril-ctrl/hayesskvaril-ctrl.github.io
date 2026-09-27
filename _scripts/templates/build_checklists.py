"""Builds tools/obligation-checklists.html and assets/templates/obligation-checklists.xlsx
from checklists_data.py. Run from the repo root, then run _scripts/sync_layout.py and recalc the xlsx.
"""
from pathlib import Path
import sys
import json
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.formatting.rule import FormulaRule

sys.path.insert(0, str(Path(__file__).parent))
from checklists_data import CHECKLISTS
from xlsx_common import (f, header_row, style_cell, list_validation, title_block, print_setup,
                         INPUT_FILL, CALC_FILL, DISCLAIMER, NAVY)

ROOT = Path(__file__).resolve().parents[2]
HTML_OUT = ROOT / "tools" / "obligation-checklists.html"
XLSX_OUT = ROOT / "assets" / "templates" / "obligation-checklists.xlsx"
REVIEWED = "27 September 2026"
BANDS = json.dumps([
    [100, "All ticked", "Every item is in place on your own assessment. Make sure each tick is backed by evidence."],
    [70, "Mostly in place", "Some gaps remain. Prioritise items that affect customers or regulator deadlines."],
    [0, "Significant gaps", "Several items aren't yet in place. Use the downloadable version to plan owners and actions."],
])

# ---------------------------------------------------------------- HTML
toc = "\n".join(f'      <li><a href="#{c["id"]}">{escape(c["title"])}</a> <span class="small">({len(c["items"])} items)</span></li>'
                for c in CHECKLISTS)
sections = []
for c in CHECKLISTS:
    items = "\n".join(
        f'      <label><input type="checkbox"> <span>{escape(t)} <span class="check-ref">{escape(ref)}</span></span></label>'
        for t, ref in c["items"])
    sections.append(f'''  <section class="checklist-section" id="{c["id"]}" aria-labelledby="{c["id"]}-h">
    <h2 id="{c["id"]}-h">{escape(c["title"])}</h2>
    <p><strong>Applies to:</strong> {escape(c["applies"])} <a class="xref" href="{c["page"]}">Read the explainer</a></p>
    <div class="widget checklist-widget" data-bands="{escape(BANDS, quote=True)}">
{items}
      <p class="result" aria-live="polite"></p>
    </div>
  </section>''')

html = f'''<!DOCTYPE html>
<html lang="en-AU">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Regulatory obligation checklists | RiskLens Australia</title>
<meta name="description" content="Free self-check checklists for key Australian regulatory obligations: CPS 230 operational risk, ASIC breach reporting, complaints handling (RG 271), design and distribution (RG 274), whistleblower programs and privacy, with a downloadable Excel version.">
<link rel="stylesheet" href="/styles.css">
</head>
<body>

<!-- HEADER:START -->
<!-- HEADER:END -->

<main id="main" class="article">
  <nav class="breadcrumb" aria-label="Breadcrumb"><ol><li><a href="/">Home</a></li><li><a href="/tools/">Tools</a></li><li>Obligation checklists</li></ol></nav>

  <h1>Regulatory obligation checklists</h1>
  <p class="summary">Quick self-checks for six sets of obligations that financial services organisations commonly manage. Tick what's in place here, or download the spreadsheet version to record evidence, owners and actions.</p>
  <div class="page-meta">
    <span class="level level-intermediate">Intermediate</span>
    <span>Interactive checklists and Excel download</span>
    <span>Last reviewed: {REVIEWED}</span>
  </div>

  <div class="download-box">
    <a class="button" href="/assets/templates/obligation-checklists.xlsx" download>Download all checklists (.xlsx)</a>
    <p class="file-meta">One tab per checklist, with status drop-downs, evidence, owner, action and due date columns, plus a summary tab. Works in Excel, Google Sheets and LibreOffice.</p>
  </div>

  <aside class="takeaways" aria-labelledby="kt">
    <h2 id="kt">Key takeaways</h2>
    <ul>
      <li>Each item is a plain-English summary of a requirement or good practice, with a reference to where it comes from. It is not the legal text.</li>
      <li>A tick should mean "we could show evidence of this", not just "we think we do this".</li>
      <li>These checklists are for learning and first-pass self-assessment. A full compliance review needs your organisation's own obligations register and the current law and guidance.</li>
      <li>Ticks on this page aren't saved. Use the spreadsheet to keep a record.</li>
    </ul>
  </aside>

  <nav class="pathway-toc" aria-label="Checklists">
    <h2>Checklists on this page</h2>
    <ul>
{toc}
    </ul>
  </nav>

{chr(10).join(sections)}

  <section class="related" aria-labelledby="related-h">
    <h2 id="related-h">Related topics</h2>
    <ul>
      <li data-href="/compliance/designing-a-compliance-program.html" data-label="Designing a compliance program"></li>
      <li data-href="/compliance/breach-reporting.html" data-label="Breach and incident reporting obligations"></li>
      <li data-href="/tools/breach-register-template.html" data-label="Breach register template"></li>
      <li data-href="/standards/" data-label="Regulatory and standards library"></li>
    </ul>
  </section>

  <section class="sources" aria-labelledby="sources-h">
    <h2 id="sources-h">Sources</h2>
    <ol>
      <li>Australian Prudential Regulation Authority, <a href="https://www.apra.gov.au/standards/cps-230">Prudential Standard CPS 230 Operational Risk Management</a>.</li>
      <li>Australian Securities and Investments Commission, <a href="https://www.asic.gov.au/regulatory-resources/financial-services/reportable-situations-for-afs-and-credit-licensees">Reportable situations for AFS and credit licensees</a> (RG 78), <a href="https://www.asic.gov.au/regulatory-resources/find-a-document/regulatory-guides/rg-271-internal-dispute-resolution">RG 271 Internal dispute resolution</a>, <a href="https://download.asic.gov.au/media/etgm1amc/rg274-published-10-september-2024.pdf">RG 274 Product design and distribution obligations</a>, and <a href="https://www.asic.gov.au/about-asic/asic-investigations-and-enforcement/whistleblowing/protections-for-corporate-sector-whistleblowers">Protections for corporate sector whistleblowers</a>.</li>
      <li>Office of the Australian Information Commissioner, <a href="https://www.oaic.gov.au/privacy/australian-privacy-principles">Australian Privacy Principles</a> and <a href="https://www.oaic.gov.au/privacy/notifiable-data-breaches/about-the-notifiable-data-breaches-scheme">About the Notifiable Data Breaches scheme</a>.</li>
      <li>Each checklist links to a RiskLens explainer page with fuller sources.</li>
    </ol>
  </section>

  <p class="last-reviewed">Last reviewed: {REVIEWED}</p>
</main>

<!-- FOOTER:START -->
<!-- FOOTER:END -->

<script src="/scripts/checklist.js"></script>
</body>
</html>
'''
HTML_OUT.write_text(html, encoding="utf-8")
print("wrote", HTML_OUT.relative_to(ROOT))

# ---------------------------------------------------------------- XLSX
TABS = {"cps-230": "CPS 230", "breach-reporting": "Breach reporting", "idr": "Complaints (IDR)",
        "ddo": "DDO", "whistleblower": "Whistleblower", "privacy": "Privacy"}
STATUSES = ["Met", "Partly met", "Not met", "Not applicable", "Not assessed"]
wb = Workbook()
sm = wb.active
sm.title = "Summary"
title_block(sm, "Regulatory obligation checklists", "RiskLens Australia · free template · version 1.0 (September 2026)")
sm["A3"] = ("How to use: on each checklist tab, choose a status for every item, note the evidence, and assign owners "
            "and actions for gaps. This summary updates automatically. Pale yellow cells are for your entries.")
sm["A3"].font = f(10, italic=True, color="4A5568")
sm["A3"].alignment = Alignment(wrap_text=True)
sm.merge_cells("A3:H3")
sm.row_dimensions[3].height = 30
header_row(sm, 5, ["Checklist", "Items", "Met", "Partly met", "Not met", "Not applicable", "Not assessed", "% met (of applicable)"],
           [44, 8, 8, 11, 9, 14, 13, 18])

lists = wb.create_sheet("Lists")
lists["A1"] = "Status"
for i, s in enumerate(STATUSES):
    lists[f"A{2 + i}"] = s

for n, c in enumerate(CHECKLISTS):
    name = TABS[c["id"]]
    ws = wb.create_sheet(name)
    title_block(ws, c["title"], "Applies to: " + c["applies"])
    ws["A3"] = f"Explainer: https://hayesskvaril-ctrl.github.io{c['page']}"
    ws["A3"].font = f(9, color="1D4ED8")
    header_row(ws, 5, ["No.", "Requirement or good practice (summary)", "Reference", "Status", "Evidence", "Owner", "Action required", "Due date"],
               [6, 60, 20, 15, 36, 16, 34, 12])
    first = 6
    for i, (text, ref) in enumerate(c["items"]):
        r = first + i
        for col, v in enumerate([i + 1, text, ref], start=1):
            style_cell(ws.cell(row=r, column=col, value=v), center=(col == 1))
        for col in range(4, 9):
            style_cell(ws.cell(row=r, column=col), INPUT_FILL)
        ws.cell(row=r, column=8).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=4, value="Not assessed")
    if n == 0:  # one worked example row, shown in green, so users see the expected format
        import datetime
        from xlsx_common import EXAMPLE_FILL
        for col, v in zip(range(4, 9), ["Partly met", "Board approved the framework (minutes 12/05/2026); roles documented in the RMS.",
                                        "Chief Risk Officer", "Add CPS 230 roles to senior managers' accountability statements.",
                                        datetime.date(2026, 12, 31)]):
            cell = ws.cell(row=first, column=col, value=v)
            cell.fill = EXAMPLE_FILL
        ws.cell(row=4, column=1, value="Green row = worked example. Replace it with your own assessment.").font = f(9, italic=True, color="15803D")
    last = first + len(c["items"]) - 1
    list_validation(ws, f"=Lists!$A$2:$A${1 + len(STATUSES)}", f"D{first}:D{last}")
    for status, colour in (("Met", "BFE5CB"), ("Partly met", "FDE68A"), ("Not met", "FCA5A5")):
        ws.conditional_formatting.add(f"D{first}:D{last}", FormulaRule(formula=[f'D{first}="{status}"'],
                                                                     fill=PatternFill("solid", fgColor=colour)))
    ws.freeze_panes = f"A{first}"
    # summary row
    sr = 6 + n
    ref = f"'{name}'!$D${first}:$D${last}"
    style_cell(sm.cell(row=sr, column=1, value=c["title"]))
    style_cell(sm.cell(row=sr, column=2, value=len(c["items"])), center=True)
    for j, s in enumerate(STATUSES):
        style_cell(sm.cell(row=sr, column=3 + j, value=f'=COUNTIF({ref},"{s}")'), CALC_FILL, center=True)
    pct = sm.cell(row=sr, column=8, value=f'=IF((B{sr}-F{sr})=0,"",C{sr}/(B{sr}-F{sr}))')
    style_cell(pct, CALC_FILL, center=True)
    pct.number_format = "0%"

end = 6 + len(CHECKLISTS)
sm.cell(row=end + 1, column=1, value=DISCLAIMER).font = f(9, italic=True, color="4A5568")
sm.cell(row=end + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
sm.merge_cells(start_row=end + 1, start_column=1, end_row=end + 1, end_column=8)
sm.row_dimensions[end + 1].height = 40
wb.move_sheet("Lists", offset=len(wb.sheetnames))
lists.sheet_state = "visible"
print_setup(wb)
wb.active = 0
wb.save(XLSX_OUT)
print("wrote", XLSX_OUT.relative_to(ROOT))
