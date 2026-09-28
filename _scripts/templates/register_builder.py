"""Generic builder for register-style spreadsheet templates (Read me, register, dashboard, lists).

A spec is a dict with:
  file, title, subtitle, readme [(heading, text)], lists {name: [values]},
  sheet (register sheet name), rows (blank rows), cols [col dicts], examples [ {header: value} ],
  dashboard [(label, formula)] where formulas can use R('Header') for the column range,
  breakdowns [(title, list name, header)]  -> COUNTIF tables on the dashboard,
  extra (optional callable(wb, helpers) to add more sheets).
Column dict: h (header), w (width), kind 'in' | 'list' | 'calc', lst (list name for drop-downs),
  formula (for calc: python format string with {r} = row and {X} = column letter X, e.g. {C}{r}),
  fmt (number format), rating (True to colour Low/Medium/High/Extreme), note (header comment).
"""
from pathlib import Path
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import (f, header_row, style_cell, list_validation, title_block, print_setup, rating_colours,
                         INPUT_FILL, CALC_FILL, EXAMPLE_FILL, DISCLAIMER, NAVY)

ROOT = Path(__file__).resolve().parents[2]


def build(spec):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    title_block(ws, spec["title"], spec["subtitle"])
    r = 4
    for k, v in spec["readme"] + [("Colour key", "Pale yellow cells: fill these in (many have drop-down lists).\nGrey-blue cells: calculated automatically; don't type over them.\nGreen rows: worked examples. Replace or delete them."),
                                  ("Important", DISCLAIMER)]:
        ws.cell(row=r, column=1, value=k).font = f(11, True, NAVY)
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        c = ws.cell(row=r, column=2, value=v)
        c.font = f(10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(30, 15 * (v.count("\n") + 1 + len(v) // 95))
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    for i, (label, fill) in enumerate([("Fill in", INPUT_FILL), ("Calculated", CALC_FILL), ("Example", EXAMPLE_FILL)]):
        style_cell(ws.cell(row=r + 1 + i, column=1, value=label), fill)

    # Lists
    names = list(spec.get("lists", {}).keys())
    if names:
        ls = wb.create_sheet("Lists")
        title_block(ls, "Drop-down lists", "Edit these values to suit your organisation. The drop-downs read from this tab.")
        for j, name in enumerate(names):
            col = 1 + j
            style_cell(ls.cell(row=4, column=col, value=name), bold=True)
            for i, v in enumerate(spec["lists"][name]):
                style_cell(ls.cell(row=5 + i, column=col, value=v), INPUT_FILL)
            ls.column_dimensions[get_column_letter(col)].width = 28

    def list_ref(name):
        j = names.index(name)
        col = get_column_letter(1 + j)
        return f"=Lists!${col}$5:${col}${4 + len(spec['lists'][name])}"

    # Register
    rg = wb.create_sheet(spec["sheet"], 1)
    title_block(rg, spec["sheet"], spec.get("sheet_note", "Fill in the yellow cells. Grey-blue cells calculate automatically."))
    HR = 4
    cols = spec["cols"]
    letters = {c["h"]: get_column_letter(i + 1) for i, c in enumerate(cols)}
    colmap = {get_column_letter(i + 1): get_column_letter(i + 1) for i in range(len(cols))}
    header_row(rg, HR, [c["h"] for c in cols], [c.get("w", 16) for c in cols], height=44)
    for i, c in enumerate(cols):
        if c.get("note"):
            rg.cell(row=HR, column=i + 1).comment = Comment(c["note"], "RiskLens")
    first, last = HR + 1, HR + spec["rows"]
    examples = spec.get("examples", [])
    for rr in range(first, last + 1):
        ex = examples[rr - first] if rr - first < len(examples) else None
        for i, c in enumerate(cols):
            cell = rg.cell(row=rr, column=i + 1)
            if c["kind"] == "calc":
                cell.value = "=" + c["formula"].format(r=rr, **colmap).lstrip("=")
                style_cell(cell, CALC_FILL, fmt=c.get("fmt"))
            else:
                if ex is not None and c["h"] in ex:
                    cell.value = ex[c["h"]]
                style_cell(cell, EXAMPLE_FILL if ex is not None else INPUT_FILL, fmt=c.get("fmt"))
            if ex is not None and c["kind"] == "calc":
                cell.fill = CALC_FILL
    for i, c in enumerate(cols):
        L = get_column_letter(i + 1)
        rngs = f"{L}{first}:{L}{last}"
        if c["kind"] == "list":
            list_validation(rg, list_ref(c["lst"]), rngs)
        if c.get("rating"):
            rating_colours(rg, rngs)
    rg.freeze_panes = rg.cell(row=first, column=2)
    rg.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"

    def R(h):
        L = letters[h]
        return f"'{spec['sheet']}'!${L}${first}:${L}${last}"

    # Dashboard
    if spec.get("dashboard"):
        db = wb.create_sheet("Dashboard", 2)
        title_block(db, "Dashboard", "Updates automatically from the register.")
        header_row(db, 4, ["Measure", "Value"], [48, 16])
        for i, (label, formula) in enumerate(spec["dashboard"]):
            style_cell(db.cell(row=5 + i, column=1, value=label))
            style_cell(db.cell(row=5 + i, column=2, value=formula(R)), CALC_FILL, center=True)
        row0 = 6 + len(spec["dashboard"])
        for title, lname, h in spec.get("breakdowns", []):
            header_row(db, row0, [title, "Count"])
            for i, v in enumerate(spec["lists"][lname]):
                style_cell(db.cell(row=row0 + 1 + i, column=1, value=v))
                style_cell(db.cell(row=row0 + 1 + i, column=2, value=f'=COUNTIF({R(h)},A{row0 + 1 + i})'), CALC_FILL, center=True)
            row0 += len(spec["lists"][lname]) + 2

    if spec.get("extra"):
        spec["extra"](wb, dict(R=R, letters=letters, first=first, last=last))
    print_setup(wb)
    out = ROOT / "assets" / "templates" / spec["file"]
    wb.save(out)
    print("wrote", out.relative_to(ROOT))
    return out
