"""Builds assets/templates/breach-register-template.xlsx.

Run from the repo root:  python3 _scripts/templates/build_breach_register.py
Then recalculate with LibreOffice (see README in this folder).
"""
from pathlib import Path
import sys
import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import (f, header_row, style_cell, list_validation, title_block, print_setup,
                         INPUT_FILL, CALC_FILL, EXAMPLE_FILL, DISCLAIMER, NAVY)

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "assets" / "templates" / "breach-register-template.xlsx"
ROWS = 100

wb = Workbook()

# ---------------------------------------------------------------- Read me
ws = wb.active
ws.title = "Read me"
title_block(ws, "Breach register template", "RiskLens Australia · free template · version 1.0 (September 2026)")
lines = [
    ("What this is", "A register for recording breaches and possible breaches of obligations, assessing whether they must be reported, tracking notification deadlines, and following remediation through to closure. Background: https://hayesskvaril-ctrl.github.io/risk-management/incident-and-breach-management.html"),
    ("How to use it", "1. Log every breach or possible breach as soon as it is identified, including the date the organisation became aware.\n"
                      "2. Assess significance and record your reasoning, especially where you decide a matter is not reportable.\n"
                      "3. For each regulator, record whether a report is needed, the deadline and the date lodged. The ASIC deadline calculates automatically.\n"
                      "4. Track root cause, corrective actions and customer remediation until closure.\n"
                      "5. The 'Dashboard' tab summarises the register for management and board reporting."),
    ("ASIC deadline", "For ASIC reportable situations, the register calculates 30 calendar days from the 'Date aware' column. That is when the licensee first knew of, or was reckless about, reasonable grounds to believe a reportable situation had arisen. Check RG 78 and any current ASIC relief for your circumstances."),
    ("Other deadlines", "APRA, OAIC, AUSTRAC and other deadlines vary by regime (from 24 hours to 30 days). Enter those deadline dates yourself. See https://hayesskvaril-ctrl.github.io/compliance/breach-reporting.html"),
    ("Colour key", "Pale yellow cells: fill these in (many have drop-down lists).\nGrey-blue cells: calculated automatically; don't type over them.\nGreen row: a worked example. Replace or delete it."),
    ("Status formulas", "Deadline status shows 'Due in X days', 'OVERDUE', 'On time' or 'Late', and uses today's date, so it updates each time the file is opened."),
    ("Important", DISCLAIMER),
]
r = 4
for k, v in lines:
    ws.cell(row=r, column=1, value=k).font = f(11, True, NAVY)
    ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
    c = ws.cell(row=r, column=2, value=v)
    c.font = f(10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = max(30, 15 * (v.count("\n") + 1 + len(v) // 95))
    r += 1
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 100
for i, (label, fill) in enumerate([("Fill in", INPUT_FILL), ("Calculated", CALC_FILL), ("Example", EXAMPLE_FILL)]):
    style_cell(ws.cell(row=r + 1 + i, column=1, value=label), fill)

# ---------------------------------------------------------------- Lists
ls = wb.create_sheet("Lists")
title_block(ls, "Drop-down lists", "Edit these values to suit your organisation. The register's drop-downs read from this tab.")
LISTS = [
    ("Source of obligation", ["Corporations Act", "SIS Act / Regulations", "National Credit Act", "ASIC Act",
                              "ASIC instrument or RG", "APRA prudential standard", "Privacy Act", "AML/CTF Act",
                              "Industry code", "Licence condition", "Contract", "Internal policy", "Other"]),
    ("How identified", ["Control / monitoring", "Complaint", "Staff report", "Internal audit", "External audit",
                        "Service provider", "Regulator", "Whistleblower", "Other"]),
    ("Yes / No", ["Yes", "No", "Under assessment", "N/A"]),
    ("Remediation status", ["Not required", "Scoping", "Calculating", "Paying", "Complete"]),
    ("Breach status", ["Open", "Under investigation", "Remediation in progress", "Closed"]),
]
for j, (name, vals) in enumerate(LISTS):
    col = 1 + j
    h = ls.cell(row=4, column=col, value=name)
    style_cell(h, bold=True)
    for i, v in enumerate(vals):
        style_cell(ls.cell(row=5 + i, column=col, value=v), INPUT_FILL)
    ls.column_dimensions[get_column_letter(col)].width = 26


def list_ref(j, n):
    col = get_column_letter(1 + j)
    return f"=Lists!${col}$5:${col}${4 + n}"


# ---------------------------------------------------------------- Register
rg = wb.create_sheet("Breach register")
title_block(rg, "Breach register", "Fill in the pale yellow cells. Deadlines, status and days open calculate automatically.")
COLS = [
    ("Breach ID", 10, "in"), ("Date identified", 12, "date"), ("Date aware (clock start)", 13, "date"),
    ("Description", 40, "in"), ("Business unit", 16, "in"), ("Obligation breached", 30, "in"),
    ("Source of obligation", 18, "in"), ("How identified", 16, "in"), ("Breach owner", 16, "in"),
    ("Customers / members affected", 12, "num"), ("Estimated financial impact ($)", 14, "money"),
    ("Significant?", 13, "in"), ("Significance reasoning", 40, "in"),
    ("ASIC report required?", 12, "in"), ("ASIC deadline (30 days)", 13, "calcdate"), ("ASIC lodged date", 12, "date"),
    ("ASIC status", 15, "calc"),
    ("APRA notification required?", 12, "in"), ("APRA deadline", 12, "date"), ("APRA notified date", 12, "date"),
    ("APRA status", 15, "calc"),
    ("OAIC data breach?", 12, "in"), ("Other notifications (AUSTRAC, customers, others)", 28, "in"),
    ("Customer remediation required?", 13, "in"), ("Remediation status", 15, "in"),
    ("Root cause", 34, "in"), ("Corrective actions", 34, "in"), ("Action owner", 15, "in"), ("Action due date", 12, "date"),
    ("Breach status", 16, "in"), ("Date closed", 12, "date"), ("Days open", 9, "calc"), ("Notes", 30, "in"),
]
HR = 4
header_row(rg, HR, [c[0] for c in COLS], [c[1] for c in COLS], height=48)
first, last = HR + 1, HR + ROWS
col_of = {name: get_column_letter(i + 1) for i, (name, _, _) in enumerate(COLS)}
C = lambda name: col_of[name]


def status_formula(req, deadline, done, r):
    return (f'=IF(OR({req}{r}<>"Yes",NOT(ISNUMBER({deadline}{r}))),"",'
            f'IF(ISNUMBER({done}{r}),IF({done}{r}<={deadline}{r},"On time","Late"),'
            f'IF(TODAY()>{deadline}{r},"OVERDUE","Due in "&({deadline}{r}-TODAY())&" days")))')


for r in range(first, last + 1):
    for i, (name, _, kind) in enumerate(COLS, start=1):
        c = rg.cell(row=r, column=i)
        style_cell(c, CALC_FILL if kind.startswith("calc") else INPUT_FILL)
        if kind in ("date", "calcdate"):
            c.number_format = "DD/MM/YYYY"
        elif kind == "money":
            c.number_format = '$#,##0;($#,##0);"-"'
        elif kind == "num":
            c.number_format = "#,##0"
    rg[f"{C('ASIC deadline (30 days)')}{r}"] = (
        f'=IF(AND({C("ASIC report required?")}{r}="Yes",ISNUMBER({C("Date aware (clock start)")}{r})),'
        f'{C("Date aware (clock start)")}{r}+30,"")')
    rg[f"{C('ASIC status')}{r}"] = status_formula(C("ASIC report required?"), C("ASIC deadline (30 days)"), C("ASIC lodged date"), r)
    rg[f"{C('APRA status')}{r}"] = status_formula(C("APRA notification required?"), C("APRA deadline"), C("APRA notified date"), r)
    aware, closed = C("Date aware (clock start)"), C("Date closed")
    rg[f"{C('Days open')}{r}"] = f'=IF(ISNUMBER({aware}{r}),IF(ISNUMBER({closed}{r}),{closed}{r}-{aware}{r},TODAY()-{aware}{r}),"")'
    rg[f"{C('Days open')}{r}"].number_format = "0"

EX = {
    "Breach ID": "B-2026-001", "Date identified": datetime.date(2026, 9, 1), "Date aware (clock start)": datetime.date(2026, 9, 1),
    "Description": "Administration fee calculated on the wrong basis for about 3,000 members since a product change 18 months ago.",
    "Business unit": "Operations", "Obligation breached": "Charging fees not permitted by the governing rules / PDS; efficiently, honestly and fairly (s 912A)",
    "Source of obligation": "Corporations Act", "How identified": "Service provider", "Breach owner": "Head of Operations",
    "Customers / members affected": 3000, "Estimated financial impact ($)": 450000, "Significant?": "Yes",
    "Significance reasoning": "Likely to cause material loss to members across a large group; treated as deemed significant.",
    "ASIC report required?": "Yes", "ASIC lodged date": datetime.date(2026, 9, 24),
    "APRA notification required?": "Yes", "APRA deadline": datetime.date(2026, 10, 1), "APRA notified date": datetime.date(2026, 9, 24),
    "OAIC data breach?": "No", "Other notifications (AUSTRAC, customers, others)": "Board Risk Committee informed 10/09/2026",
    "Customer remediation required?": "Yes", "Remediation status": "Calculating",
    "Root cause": "Product change testing did not include fee calculation scenarios.",
    "Corrective actions": "Fix calculation; add fee scenarios to change testing; monthly independent fee reconciliation.",
    "Action owner": "Product Manager", "Action due date": datetime.date(2026, 12, 31), "Breach status": "Remediation in progress",
    "Notes": "One report to APRA containing all ASIC-required information can count as lodgement with ASIC for RSE licensees.",
}
for name, v in EX.items():
    rg[f"{C(name)}{first}"] = v
for i, (name, _, kind) in enumerate(COLS, start=1):
    if not kind.startswith("calc"):
        rg.cell(row=first, column=i).fill = EXAMPLE_FILL
rg.row_dimensions[first].height = 90

rng = lambda name: f"{C(name)}{first}:{C(name)}{last}"
list_validation(rg, list_ref(0, len(LISTS[0][1])), rng("Source of obligation"))
list_validation(rg, list_ref(1, len(LISTS[1][1])), rng("How identified"))
for name in ("Significant?", "ASIC report required?", "APRA notification required?", "OAIC data breach?", "Customer remediation required?"):
    list_validation(rg, list_ref(2, len(LISTS[2][1])), rng(name))
list_validation(rg, list_ref(3, len(LISTS[3][1])), rng("Remediation status"))
list_validation(rg, list_ref(4, len(LISTS[4][1])), rng("Breach status"))
dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
dv.error = "Enter a date (DD/MM/YYYY)."; dv.showErrorMessage = True
rg.add_data_validation(dv)
for name, _, kind in COLS:
    if kind == "date":
        dv.add(rng(name))

RED, GREEN, AMBER = "FCA5A5", "BFE5CB", "FDE68A"
for name in ("ASIC status", "APRA status"):
    col = C(name)
    rg.conditional_formatting.add(rng(name), FormulaRule(formula=[f'OR({col}{first}="OVERDUE",{col}{first}="Late")'],
                                                         fill=PatternFill("solid", fgColor=RED), font=f(10, True)))
    rg.conditional_formatting.add(rng(name), FormulaRule(formula=[f'{col}{first}="On time"'], fill=PatternFill("solid", fgColor=GREEN)))
    rg.conditional_formatting.add(rng(name), FormulaRule(formula=[f'LEFT({col}{first},6)="Due in"'], fill=PatternFill("solid", fgColor=AMBER)))
rg.freeze_panes = rg.cell(row=first, column=2)
rg.auto_filter.ref = f"A{HR}:{get_column_letter(len(COLS))}{last}"

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard")
title_block(db, "Breach register dashboard", "Updates automatically from the Breach register tab. Uses today's date for overdue checks.")
REG = "'Breach register'"
R = lambda name: f"{REG}!${C(name)}${first}:${C(name)}${last}"
header_row(db, 4, ["Measure", "Count"], [44, 12])
measures = [
    ("Breaches recorded", f"=COUNTA({R('Breach ID')})"),
    ("Open (not closed)", f"=COUNTA({R('Breach ID')})-COUNTIF({R('Breach status')},\"Closed\")"),
    ("Assessed as significant", f"=COUNTIF({R('Significant?')},\"Yes\")"),
    ("Still under assessment", f"=COUNTIF({R('Significant?')},\"Under assessment\")"),
    ("ASIC reports required", f"=COUNTIF({R('ASIC report required?')},\"Yes\")"),
    ("ASIC reports overdue (not yet lodged)", f"=COUNTIF({R('ASIC status')},\"OVERDUE\")"),
    ("ASIC reports lodged late", f"=COUNTIF({R('ASIC status')},\"Late\")"),
    ("APRA notifications overdue (not yet made)", f"=COUNTIF({R('APRA status')},\"OVERDUE\")"),
    ("APRA notifications made late", f"=COUNTIF({R('APRA status')},\"Late\")"),
    ("Customer remediation required", f"=COUNTIF({R('Customer remediation required?')},\"Yes\")"),
    ("Customers / members affected (total)", f"=SUM({R('Customers / members affected')})"),
    ("Estimated financial impact (total $)", f"=SUM({R('Estimated financial impact ($)')})"),
    ("Average days open (open breaches)", f"=IFERROR(AVERAGEIFS({R('Days open')},{R('Breach status')},\"<>Closed\"),0)"),
]
for i, (k, v) in enumerate(measures):
    style_cell(db.cell(row=5 + i, column=1, value=k))
    c = db.cell(row=5 + i, column=2, value=v)
    style_cell(c, CALC_FILL, center=True)
    c.number_format = '$#,##0' if "$" in k else "0"
row0 = 5 + len(measures) + 2
header_row(db, row0, ["Breaches by source of obligation", "Count"], height=24)
for i, v in enumerate(LISTS[0][1]):
    style_cell(db.cell(row=row0 + 1 + i, column=1, value=f"=Lists!$A${5 + i}"))
    c = db.cell(row=row0 + 1 + i, column=2, value=f'=COUNTIF({R("Source of obligation")},A{row0 + 1 + i})')
    style_cell(c, CALC_FILL, center=True)
row1 = row0 + len(LISTS[0][1]) + 3
header_row(db, row1, ["Breaches identified by month (last 12 months)", "Count"], height=24)
for i in range(12):
    rr = row1 + 1 + i
    m = db.cell(row=rr, column=1, value=f"=DATE(YEAR(TODAY()),MONTH(TODAY())-{11 - i},1)")
    style_cell(m); m.number_format = "MMM YYYY"; m.alignment = Alignment(horizontal="left")
    c = db.cell(row=rr, column=2, value=f'=COUNTIFS({R("Date identified")},">="&A{rr},{R("Date identified")},"<"&EDATE(A{rr},1))')
    style_cell(c, CALC_FILL, center=True)

print_setup(wb)
wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT.relative_to(ROOT))
