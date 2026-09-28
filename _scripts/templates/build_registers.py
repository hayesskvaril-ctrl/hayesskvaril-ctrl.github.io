"""Builds five register-style templates using register_builder.py:
obligations register, RCSA, KRI library, material service provider register, remediation program tracker.
Run from the repo root:  python3 _scripts/templates/build_registers.py   (then recalculate with LibreOffice; see README)
"""
from pathlib import Path
import sys, datetime
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Alignment
sys.path.insert(0, str(Path(__file__).parent))
from register_builder import build
from xlsx_common import f, style_cell, header_row, title_block, CALC_FILL

V = "version 1.0 (September 2026)"
SITE = "https://hayesskvaril-ctrl.github.io"
D = datetime.date
LEVELS = ["Low", "Medium", "High", "Extreme"]
YN = ["Yes", "No", "In progress", "N/A"]
BAND = 'IF({s}="","",IF({s}>=15,"Extreme",IF({s}>=10,"High",IF({s}>=5,"Medium","Low"))))'

def rag(ws, rng):
    first = rng.split(":")[0]
    for name, col in (("Green", "BFE5CB"), ("Amber", "FDE68A"), ("Red", "FCA5A5"), ("OVERDUE", "FCA5A5"), ("Due soon", "FDE68A")):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{name}"'], fill=PatternFill("solid", fgColor=col)))

def rag_extra(sheet, header):
    def fn(wb, h):
        ws = wb[sheet]
        L = h["letters"][header]
        rag(ws, f"{L}{h['first']}:{L}{h['last']}")
    return fn

# ------------------------------------------------------------------ 1. Obligations register
build(dict(
 file="obligations-register-template.xlsx", title="Obligations register template", subtitle=f"RiskLens Australia · free template · {V}",
 readme=[("What this is", f"A register of the organisation's compliance obligations: where each comes from, who owns it, the controls that ensure it is met, and when it was last reviewed. Background: {SITE}/compliance/designing-a-compliance-program.html"),
         ("How to use it", "1. Record each obligation in plain English, with its source reference (legislation section, standard paragraph, code clause).\n2. Assign an obligation owner and list the key controls.\n3. Rate inherent compliance risk and control effectiveness. Residual risk calculates using a simple rule.\n4. Set a review frequency. The next review date and review status calculate automatically.\n5. Use the Dashboard to report coverage, overdue reviews and higher-risk obligations."),
         ("Residual risk rule", "A simple, illustrative rule: effective controls reduce the inherent rating by one level; ineffective controls increase it by one level; partially effective or untested controls leave it unchanged. Replace it with your own methodology if you have one."),
         ("Keeping it current", "Link the register to your regulatory change process so new or changed obligations are added promptly.")],
 lists={"Source type": ["Legislation", "Regulations", "Prudential standard", "ASIC legislative instrument", "Regulatory guide (guidance)", "Industry code", "Licence condition", "Contract", "Internal policy"],
        "Rating": LEVELS, "Control effectiveness": ["Effective", "Partially effective", "Ineffective", "Not tested"]},
 sheet="Obligations register", rows=150,
 cols=[dict(h="Obligation ID", w=12, kind="in"), dict(h="Source type", w=20, kind="list", lst="Source type"),
       dict(h="Source reference", w=24, kind="in", note="E.g. Corporations Act s 912D; CPS 230 paragraph 33"),
       dict(h="Obligation (plain English)", w=44, kind="in"), dict(h="Business area", w=18, kind="in"), dict(h="Obligation owner", w=18, kind="in"),
       dict(h="Key controls", w=36, kind="in"), dict(h="Control owner", w=18, kind="in"),
       dict(h="Inherent compliance risk", w=14, kind="list", lst="Rating", rating=True),
       dict(h="Control effectiveness", w=16, kind="list", lst="Control effectiveness"),
       dict(h="Residual compliance risk", w=14, kind="calc", rating=True,
            formula='IF(OR({I}{r}="",{J}{r}=""),"",CHOOSE(MAX(1,MIN(4,MATCH({I}{r},{{"Low","Medium","High","Extreme"}},0)+IF({J}{r}="Effective",-1,IF({J}{r}="Ineffective",1,0)))),"Low","Medium","High","Extreme"))'),
       dict(h="Last reviewed", w=13, kind="in", fmt="DD/MM/YYYY"), dict(h="Review frequency (months)", w=12, kind="in"),
       dict(h="Next review due", w=13, kind="calc", fmt="DD/MM/YYYY", formula='IF(OR({L}{r}="",{M}{r}=""),"",EDATE({L}{r},{M}{r}))'),
       dict(h="Review status", w=12, kind="calc", formula='IF({N}{r}="","",IF({N}{r}<TODAY(),"OVERDUE",IF({N}{r}-TODAY()<=30,"Due soon","Current")))'),
       dict(h="Notes and recent changes", w=34, kind="in")],
 examples=[{"Obligation ID": "OB-001", "Source type": "Legislation", "Source reference": "Corporations Act ss 912D, 912DAA", "Obligation (plain English)": "Report reportable situations (including significant breaches) to ASIC within 30 days of reasonable grounds.", "Business area": "Compliance", "Obligation owner": "Head of Compliance", "Key controls": "Breach register; significance assessment procedure; 30-day deadline tracking; monthly review by Compliance Committee", "Control owner": "Compliance Manager", "Inherent compliance risk": "High", "Control effectiveness": "Effective", "Last reviewed": D(2026, 3, 31), "Review frequency (months)": 12, "Notes and recent changes": "ASIC relief extended the investigation period to 60 days from 27 June 2025."},
           {"Obligation ID": "OB-002", "Source type": "Legislation", "Source reference": "Privacy Act 1988, APP 11", "Obligation (plain English)": "Take reasonable steps to protect personal information from misuse, interference, loss and unauthorised access; destroy or de-identify it when no longer needed.", "Business area": "Technology and operations", "Obligation owner": "Chief Information Security Officer", "Key controls": "Access management; MFA; data retention schedule; annual destruction review", "Control owner": "IT Security Manager", "Inherent compliance risk": "Extreme", "Control effectiveness": "Partially effective", "Last reviewed": D(2025, 8, 15), "Review frequency (months)": 12},
           {"Obligation ID": "OB-003", "Source type": "Prudential standard", "Source reference": "CPS 230 (notification of material operational risk incidents)", "Obligation (plain English)": "Notify APRA as soon as possible, and within 72 hours, of a material operational risk incident.", "Business area": "Risk", "Obligation owner": "Chief Risk Officer", "Key controls": "Incident management procedure; materiality criteria; APRA notification checklist", "Control owner": "Operational Risk Manager", "Inherent compliance risk": "High", "Control effectiveness": "Not tested", "Last reviewed": D(2026, 7, 1), "Review frequency (months)": 6}],
 dashboard=[("Obligations recorded", lambda R: f"=COUNTA({R('Obligation ID')})"),
            ("Reviews overdue", lambda R: f"=COUNTIF({R('Review status')},\"OVERDUE\")"),
            ("Reviews due in the next 30 days", lambda R: f"=COUNTIF({R('Review status')},\"Due soon\")"),
            ("Residual Extreme", lambda R: f"=COUNTIF({R('Residual compliance risk')},\"Extreme\")"),
            ("Residual High", lambda R: f"=COUNTIF({R('Residual compliance risk')},\"High\")"),
            ("Controls not tested", lambda R: f"=COUNTIF({R('Control effectiveness')},\"Not tested\")"),
            ("Controls ineffective", lambda R: f"=COUNTIF({R('Control effectiveness')},\"Ineffective\")")],
 breakdowns=[("Obligations by source type", "Source type", "Source type"), ("Obligations by residual risk", "Rating", "Residual compliance risk")],
 extra=rag_extra("Obligations register", "Review status")))

# ------------------------------------------------------------------ 2. RCSA
build(dict(
 file="rcsa-template.xlsx", title="Risk and control self-assessment (RCSA) template", subtitle=f"RiskLens Australia · free template · {V}",
 readme=[("What this is", f"A template for business units to assess their key risks and controls, rate inherent and residual risk, and agree actions, with second-line challenge recorded. Background: {SITE}/risk-management/risk-assessment-methodologies.html and {SITE}/risk-management/control-design-and-testing.html"),
         ("How to use it", "1. List each key risk as cause, event and consequence.\n2. Rate inherent likelihood and consequence (1 to 5), before controls.\n3. List key controls and rate their design and operating effectiveness. The overall control rating takes the weaker of the two.\n4. Rate residual likelihood and consequence, after controls.\n5. Record whether the residual risk is within appetite, the actions needed and second-line challenge."),
         ("Ratings", "Score = likelihood x consequence (1 to 25). Bands: 1 to 4 Low, 5 to 9 Medium, 10 to 14 High, 15 to 25 Extreme, the same as the RiskLens risk assessment template. Change the formulas if your organisation uses different bands."),
         ("Good practice", "Self-assessments are only as good as the challenge they receive. Record the second line's challenge and any changes it led to, and test key controls rather than relying on self-assessed effectiveness.")],
 lists={"Risk category": ["Operational", "Compliance", "Conduct", "Technology and cyber", "Third party", "People", "Financial", "Strategic", "Investment", "Reputational"],
        "1 to 5": [1, 2, 3, 4, 5], "Effectiveness": ["Effective", "Partially effective", "Ineffective", "Not tested"],
        "Control type": ["Preventive", "Detective", "Corrective"], "Yes / No": ["Yes", "No"]},
 sheet="RCSA", rows=80,
 cols=[dict(h="Risk ID", w=10, kind="in"), dict(h="Process or business unit", w=18, kind="in"),
       dict(h="Risk description (cause, event, consequence)", w=44, kind="in"), dict(h="Risk category", w=16, kind="list", lst="Risk category"),
       dict(h="Risk owner", w=16, kind="in"),
       dict(h="Inherent likelihood (1-5)", w=11, kind="list", lst="1 to 5"), dict(h="Inherent consequence (1-5)", w=11, kind="list", lst="1 to 5"),
       dict(h="Inherent score", w=9, kind="calc", formula='IF(OR({F}{r}="",{G}{r}=""),"",{F}{r}*{G}{r})'),
       dict(h="Inherent rating", w=11, kind="calc", rating=True, formula=BAND.format(s="{H}{r}")),
       dict(h="Key controls", w=36, kind="in"), dict(h="Control type", w=12, kind="list", lst="Control type"), dict(h="Control owner", w=16, kind="in"),
       dict(h="Design effectiveness", w=14, kind="list", lst="Effectiveness"), dict(h="Operating effectiveness", w=14, kind="list", lst="Effectiveness"),
       dict(h="Overall control rating", w=14, kind="calc",
            formula='IF(OR({M}{r}="",{N}{r}=""),"",IF(OR({M}{r}="Ineffective",{N}{r}="Ineffective"),"Ineffective",IF(OR({M}{r}="Not tested",{N}{r}="Not tested"),"Not tested",IF(OR({M}{r}="Partially effective",{N}{r}="Partially effective"),"Partially effective","Effective"))))'),
       dict(h="Residual likelihood (1-5)", w=11, kind="list", lst="1 to 5"), dict(h="Residual consequence (1-5)", w=11, kind="list", lst="1 to 5"),
       dict(h="Residual score", w=9, kind="calc", formula='IF(OR({P}{r}="",{Q}{r}=""),"",{P}{r}*{Q}{r})'),
       dict(h="Residual rating", w=11, kind="calc", rating=True, formula=BAND.format(s="{R}{r}")),
       dict(h="Within appetite?", w=10, kind="list", lst="Yes / No"), dict(h="Actions required", w=30, kind="in"), dict(h="Action owner", w=14, kind="in"),
       dict(h="Action due date", w=12, kind="in", fmt="DD/MM/YYYY"), dict(h="Second-line challenge", w=30, kind="in"), dict(h="Assessment date", w=12, kind="in", fmt="DD/MM/YYYY")],
 examples=[{"Risk ID": "R-01", "Process or business unit": "Benefit payments", "Risk description (cause, event, consequence)": "Cause: manual keying of bank details. Event: payment made to the wrong account. Consequence: member loss, remediation, complaints and possible fraud.", "Risk category": "Operational", "Risk owner": "Head of Member Services", "Inherent likelihood (1-5)": 4, "Inherent consequence (1-5)": 4, "Key controls": "Call-back verification of changed bank details; four-eyes approval over $10,000; daily exception report", "Control type": "Preventive", "Control owner": "Payments Team Leader", "Design effectiveness": "Effective", "Operating effectiveness": "Partially effective", "Residual likelihood (1-5)": 3, "Residual consequence (1-5)": 3, "Within appetite?": "No", "Actions required": "Automate bank detail verification; retrain staff on call-back procedure", "Action owner": "Head of Member Services", "Action due date": D(2026, 12, 31), "Second-line challenge": "Operational Risk challenged the operating rating after three call-back exceptions in Q2 testing; residual likelihood raised from 2 to 3.", "Assessment date": D(2026, 9, 1)},
           {"Risk ID": "R-02", "Process or business unit": "Member communications", "Risk description (cause, event, consequence)": "Cause: marketing content not reviewed by Compliance. Event: misleading statement about returns. Consequence: breach of s 1041H, deemed significant breach, remediation.", "Risk category": "Compliance", "Risk owner": "Head of Marketing", "Inherent likelihood (1-5)": 3, "Inherent consequence (1-5)": 4, "Key controls": "Mandatory compliance sign-off in the content workflow; data feed checks for performance figures", "Control type": "Preventive", "Control owner": "Compliance Manager", "Design effectiveness": "Effective", "Operating effectiveness": "Effective", "Residual likelihood (1-5)": 1, "Residual consequence (1-5)": 4, "Within appetite?": "Yes", "Second-line challenge": "Compliance confirmed via sample of 25 campaigns; no exceptions.", "Assessment date": D(2026, 9, 1)}],
 dashboard=[("Risks assessed", lambda R: f"=COUNTA({R('Risk ID')})"),
            ("Residual Extreme", lambda R: f"=COUNTIF({R('Residual rating')},\"Extreme\")"),
            ("Residual High", lambda R: f"=COUNTIF({R('Residual rating')},\"High\")"),
            ("Outside appetite", lambda R: f"=COUNTIF({R('Within appetite?')},\"No\")"),
            ("Controls rated ineffective", lambda R: f"=COUNTIF({R('Overall control rating')},\"Ineffective\")"),
            ("Controls not tested", lambda R: f"=COUNTIF({R('Overall control rating')},\"Not tested\")")],
 breakdowns=[("Risks by category", "Risk category", "Risk category")]))

# ------------------------------------------------------------------ 3. KRI library
kris = [
 ("Operational", "Number of operational incidents raised", "Volume of incidents recorded in the period", "Lagging", "Count", "Monthly", 10, 20, "Higher is worse", "A sudden fall can mean under-reporting, not improvement."),
 ("Operational", "Processing backlog older than 5 business days", "Transactions not processed within service standards", "Leading", "Count", "Weekly", 50, 200, "Higher is worse", "Link thresholds to member harm and legal timeframes (e.g. rollovers)."),
 ("Operational", "Manual workarounds in critical processes", "Number of manual workarounds in place for critical operations", "Leading", "Count", "Quarterly", 3, 6, "Higher is worse", ""),
 ("Compliance", "Breaches identified in the period", "New breaches or possible breaches recorded", "Lagging", "Count", "Monthly", 5, 10, "Higher is worse", ""),
 ("Compliance", "Regulatory reports lodged late", "Reports to ASIC, APRA, OAIC or AUSTRAC lodged after the deadline", "Lagging", "Count", "Monthly", 0, 1, "Higher is worse", "Any late lodgement should usually be red."),
 ("Compliance", "Overdue compliance actions", "Compliance issues past their agreed due date", "Leading", "Count", "Monthly", 3, 8, "Higher is worse", ""),
 ("Conduct", "Complaints per 1,000 members", "Complaint rate, normalised for size", "Lagging", "Rate", "Monthly", 2, 4, "Higher is worse", "Normalise so growth doesn't look like deterioration."),
 ("Conduct", "Complaints resolved within RG 271 timeframes", "Share of complaints closed within the required maximum", "Lagging", "%", "Monthly", 95, 90, "Lower is worse", ""),
 ("Conduct", "AFCA complaints received", "External disputes lodged with AFCA", "Lagging", "Count", "Quarterly", 10, 20, "Higher is worse", ""),
 ("Technology and cyber", "Critical patches overdue", "Critical security patches not applied within the target timeframe", "Leading", "Count", "Weekly", 0, 3, "Higher is worse", "Link to Essential Eight target maturity."),
 ("Technology and cyber", "Phishing simulation click rate", "Share of staff clicking a simulated phishing link", "Leading", "%", "Quarterly", 5, 10, "Higher is worse", ""),
 ("Technology and cyber", "Critical system availability", "Uptime of systems supporting critical operations", "Lagging", "%", "Monthly", 99.5, 99, "Lower is worse", "Compare with CPS 230 tolerance levels."),
 ("Third party", "Material service providers with overdue reviews", "Providers whose annual due diligence or performance review is overdue", "Leading", "Count", "Quarterly", 0, 2, "Higher is worse", ""),
 ("Third party", "Service level breaches by material providers", "SLA misses in the period", "Lagging", "Count", "Monthly", 2, 5, "Higher is worse", ""),
 ("People", "Unplanned staff turnover in key teams", "Annualised voluntary turnover in critical roles", "Leading", "%", "Quarterly", 15, 25, "Higher is worse", ""),
 ("People", "Mandatory training overdue", "Staff with overdue compliance training", "Leading", "%", "Monthly", 5, 10, "Higher is worse", ""),
 ("Investment (super)", "Options below performance objective", "Investment options behind their return objective over the objective period", "Lagging", "Count", "Quarterly", 1, 3, "Higher is worse", ""),
 ("Liquidity (super)", "Liquid assets as % of option assets", "Liquidity buffer against switching and rollovers", "Leading", "%", "Monthly", 10, 7, "Lower is worse", "Set from stress test results, not historical averages."),
 ("Fraud", "Bank detail changes followed by payment within 48 hours", "A common account takeover pattern", "Leading", "Count", "Weekly", 5, 15, "Higher is worse", ""),
 ("Remediation", "Remediation programs past target date", "Open remediation programs behind plan", "Leading", "Count", "Monthly", 0, 2, "Higher is worse", ""),
]
DEMO = {0: 8, 1: 120, 2: 7}  # example current values to show Green, Amber and Red
def kri_notes(wb, h):
    ws = wb.create_sheet("Threshold design notes")
    title_block(ws, "Designing KRI thresholds", "Notes to help set green, amber and red thresholds.")
    notes = [("Link to appetite", "Red should mean 'outside appetite' and trigger defined escalation. Amber is an early warning with room to act."),
             ("Direction", "Some KRIs are worse when higher (incidents), others when lower (service levels). Set the Direction column so the status calculates correctly."),
             ("Evidence", "Base thresholds on history, peer data, tolerance levels and stress testing. Record the rationale in the design notes column."),
             ("Leading and lagging", "Balance lagging indicators (what happened) with leading ones (what might happen). Leading indicators give time to act."),
             ("Test predictive value", "Check whether indicators actually moved before past incidents. Replace ones that didn't."),
             ("Review", "Review KRIs and thresholds at least annually, and after major incidents or changes to the business.")]
    for i, (k, v) in enumerate(notes):
        ws.cell(row=4 + i, column=1, value=k).font = f(11, True, "0F2942")
        c = ws.cell(row=4 + i, column=2, value=v); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4 + i].height = 32
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 100
    rag(wb["KRI library"], f"{h['letters']['Status']}{h['first']}:{h['letters']['Status']}{h['last']}")
build(dict(
 file="kri-library.xlsx", title="Key risk indicator (KRI) library", subtitle=f"RiskLens Australia · free template · {V}",
 readme=[("What this is", f"A library of 20 example key risk indicators across risk types, with threshold design notes, and space to record current values (the first three have example values). Status (Green, Amber, Red) calculates automatically. Background: {SITE}/governance/board-risk-reporting.html"),
         ("How to use it", "1. Choose the KRIs relevant to your risk profile and delete the rest.\n2. Set green and red thresholds from your risk appetite. Values between them are amber.\n3. Set the direction (higher is worse, or lower is worse).\n4. Enter the current value each period. Status calculates automatically.\n5. See 'Threshold design notes' for guidance."),
         ("Example thresholds", "All thresholds in the examples are illustrative only. They are not benchmarks. Set your own from your risk appetite statement and data.")],
 lists={"Risk type": ["Operational", "Compliance", "Conduct", "Technology and cyber", "Third party", "People", "Investment (super)", "Liquidity (super)", "Fraud", "Remediation", "Other"],
        "Leading or lagging": ["Leading", "Lagging"], "Frequency": ["Daily", "Weekly", "Monthly", "Quarterly", "Annually"], "Direction": ["Higher is worse", "Lower is worse"]},
 sheet="KRI library", rows=60,
 cols=[dict(h="KRI ID", w=9, kind="in"), dict(h="Risk type", w=16, kind="list", lst="Risk type"), dict(h="KRI", w=36, kind="in"),
       dict(h="What it measures", w=38, kind="in"), dict(h="Leading or lagging", w=11, kind="list", lst="Leading or lagging"),
       dict(h="Unit", w=8, kind="in"), dict(h="Frequency", w=11, kind="list", lst="Frequency"),
       dict(h="Green up to (or from)", w=11, kind="in", note="The limit of the green zone. For 'Lower is worse' KRIs, values at or above this are green."),
       dict(h="Red at (or below)", w=11, kind="in", note="The red trigger. For 'Higher is worse' KRIs, values at or above this are red; for 'Lower is worse', values at or below this are red."),
       dict(h="Direction", w=14, kind="list", lst="Direction"), dict(h="Current value", w=10, kind="in"),
       dict(h="Status", w=10, kind="calc", formula='IF(OR({K}{r}="",{H}{r}="",{I}{r}="",{J}{r}=""),"",IF({J}{r}="Higher is worse",IF({K}{r}>={I}{r},"Red",IF({K}{r}>{H}{r},"Amber","Green")),IF({K}{r}<={I}{r},"Red",IF({K}{r}<{H}{r},"Amber","Green"))))'),
       dict(h="Owner", w=16, kind="in"), dict(h="Design notes", w=40, kind="in")],
 examples=[{"KRI ID": f"KRI-{i+1:02d}", "Risk type": k[0], "KRI": k[1], "What it measures": k[2], "Leading or lagging": k[3], "Unit": k[4], "Frequency": k[5],
            "Green up to (or from)": k[6], "Red at (or below)": k[7], "Direction": k[8], "Design notes": k[9], **({"Current value": DEMO[i]} if i in DEMO else {})} for i, k in enumerate(kris)],
 dashboard=[("KRIs with a current value", lambda R: f"=COUNT({R('Current value')})"),
            ("Red", lambda R: f"=COUNTIF({R('Status')},\"Red\")"), ("Amber", lambda R: f"=COUNTIF({R('Status')},\"Amber\")"), ("Green", lambda R: f"=COUNTIF({R('Status')},\"Green\")")],
 breakdowns=[("KRIs by risk type", "Risk type", "Risk type")],
 extra=kri_notes))

# ------------------------------------------------------------------ 4. Material service provider register
build(dict(
 file="material-service-provider-register.xlsx", title="Material service provider register", subtitle=f"RiskLens Australia · free template · {V}",
 readme=[("What this is", f"An internal register of material service providers and the critical operations they support, for managing arrangements under CPS 230. Background: {SITE}/risk-management/third-party-risk.html and {SITE}/standards/cps-230.html"),
         ("APRA submission", "APRA publishes its own material service provider register template (updated 30 April 2026), which is its preferred way to submit registers. Use APRA's official template for submission. This workbook captures similar information for internal management and adds due diligence, contract and review tracking."),
         ("How to use it", "1. Record each material arrangement and the critical operations it supports.\n2. Record where services are delivered and data is held, including offshore locations and key fourth parties.\n3. Track contract dates, due diligence and performance reviews. Alerts calculate automatically.\n4. Record APRA notifications for new or materially changed arrangements."),
         ("Alerts", "Contract alert: 'Ending soon' within 180 days of contract end. Review alert: 'OVERDUE' if the last review was more than 12 months ago.")],
 lists={"Yes / No": ["Yes", "No", "Unknown"], "Service category": ["Core technology", "Cloud hosting", "Fund administration", "Custody", "Investment management", "Insurance (group insurer)", "Claims management", "Credit assessment", "Payments", "Internal audit", "Risk management", "Data centre", "Other"],
        "Performance": ["Meeting expectations", "Some concerns", "Serious concerns"]},
 sheet="MSP register", rows=60,
 cols=[dict(h="Provider ID", w=10, kind="in"), dict(h="Provider legal name", w=26, kind="in"), dict(h="ABN or overseas identifier", w=16, kind="in"),
       dict(h="Related party?", w=10, kind="list", lst="Yes / No"), dict(h="Service category", w=18, kind="list", lst="Service category"),
       dict(h="Description of services", w=36, kind="in"), dict(h="Critical operations supported", w=28, kind="in"),
       dict(h="Service delivery locations", w=20, kind="in"), dict(h="Data storage and processing locations", w=22, kind="in"),
       dict(h="Offshore?", w=10, kind="list", lst="Yes / No"), dict(h="Key fourth parties", w=24, kind="in"),
       dict(h="Contract start", w=12, kind="in", fmt="DD/MM/YYYY"), dict(h="Contract end", w=12, kind="in", fmt="DD/MM/YYYY"),
       dict(h="Contract alert", w=12, kind="calc", formula='IF({M}{r}="","",IF({M}{r}<TODAY(),"Expired",IF({M}{r}-TODAY()<=180,"Ending soon","OK")))'),
       dict(h="Contract meets CPS 230 requirements?", w=14, kind="list", lst="Yes / No"), dict(h="Exit or contingency plan in place?", w=14, kind="list", lst="Yes / No"),
       dict(h="Last due diligence or review", w=13, kind="in", fmt="DD/MM/YYYY"),
       dict(h="Review alert", w=11, kind="calc", formula='IF({Q}{r}="","",IF(EDATE({Q}{r},12)<TODAY(),"OVERDUE",IF(EDATE({Q}{r},12)-TODAY()<=30,"Due soon","Current")))'),
       dict(h="Performance", w=16, kind="list", lst="Performance"), dict(h="APRA notified (date)", w=13, kind="in", fmt="DD/MM/YYYY"),
       dict(h="Relationship owner", w=18, kind="in"), dict(h="Notes", w=30, kind="in")],
 examples=[{"Provider ID": "MSP-01", "Provider legal name": "Example Administration Services Pty Ltd (fictional)", "ABN or overseas identifier": "00 000 000 000", "Related party?": "No", "Service category": "Fund administration", "Description of services": "Member administration, contributions processing, benefit payments and call centre", "Critical operations supported": "Fund administration; paying benefits; member enquiries", "Service delivery locations": "Australia; Philippines (call centre overflow)", "Data storage and processing locations": "Australia (cloud region: Sydney)", "Offshore?": "Yes", "Key fourth parties": "Cloud hosting provider; print and mail house", "Contract start": D(2022, 7, 1), "Contract end": D(2027, 3, 31), "Contract meets CPS 230 requirements?": "Yes", "Exit or contingency plan in place?": "Yes", "Last due diligence or review": D(2025, 8, 1), "Performance": "Some concerns", "APRA notified (date)": D(2025, 7, 10), "Relationship owner": "Head of Operations", "Notes": "Two outages this year; recovery within tolerance but close to the limit."}],
 dashboard=[("Arrangements recorded", lambda R: f"=COUNTA({R('Provider ID')})"),
            ("Offshore arrangements", lambda R: f"=COUNTIF({R('Offshore?')},\"Yes\")"),
            ("Contracts ending within 180 days", lambda R: f"=COUNTIF({R('Contract alert')},\"Ending soon\")"),
            ("Reviews overdue", lambda R: f"=COUNTIF({R('Review alert')},\"OVERDUE\")"),
            ("Contracts not yet meeting CPS 230", lambda R: f"=COUNTIF({R('Contract meets CPS 230 requirements?')},\"No\")"),
            ("No exit or contingency plan", lambda R: f"=COUNTIF({R('Exit or contingency plan in place?')},\"No\")"),
            ("Serious performance concerns", lambda R: f"=COUNTIF({R('Performance')},\"Serious concerns\")")],
 breakdowns=[("Arrangements by service category", "Service category", "Service category")],
 extra=lambda wb, h: (rag(wb["MSP register"], f"{h['letters']['Review alert']}{h['first']}:{h['letters']['Review alert']}{h['last']}"),
                      rag(wb["MSP register"], f"{h['letters']['Contract alert']}{h['first']}:{h['letters']['Contract alert']}{h['last']}"))))

# ------------------------------------------------------------------ 5. Remediation program tracker
build(dict(
 file="remediation-program-tracker.xlsx", title="Remediation program tracker", subtitle=f"RiskLens Australia · free template · {V}",
 readme=[("What this is", f"A tracker for customer remediation programs, from identifying the affected population to paying compensation, handling residual amounts and closing the program. Background: {SITE}/standards/asic-rg-277.html"),
         ("How to use it", "1. Create one row per remediation program and link it to the breach or incident.\n2. Record the affected population, the period, the method and beneficial assumptions used.\n3. Update numbers paid and amounts paid as the program progresses. Percentages and alerts calculate automatically.\n4. Record how residual (unclaimed) amounts will be handled, assurance over calculations and payments, and whether the root cause is fixed."),
         ("RG 277", "ASIC's RG 277 expects remediation to be fair, timely and efficient, with beneficial assumptions where data is missing, compensation for the time value of money, reasonable endeavours to return money, and specific options for residual amounts. Check RG 277 for the details."),
         ("Worked example", "The example row is fictional and follows the fee error case used on the RiskLens breach significance page.")],
 lists={"Yes / No": ["Yes", "No", "In progress", "N/A"], "Status": ["Scoping", "Calculating", "Approved", "Paying", "Residual handling", "Closed"],
        "Residual approach": ["Not yet decided", "Lodge as unclaimed money (ASIC, or ATO for super)", "Pay to charity or not-for-profit (where permitted)", "Other (document basis)", "Not applicable"]},
 sheet="Remediation tracker", rows=40,
 cols=[dict(h="Program ID", w=10, kind="in"), dict(h="Linked breach or incident", w=14, kind="in"), dict(h="Issue description", w=34, kind="in"),
       dict(h="Date identified", w=12, kind="in", fmt="DD/MM/YYYY"), dict(h="Period affected", w=18, kind="in"),
       dict(h="Customers affected", w=11, kind="in", fmt="#,##0"), dict(h="Method and beneficial assumptions", w=36, kind="in"),
       dict(h="Estimated compensation incl. interest ($)", w=15, kind="in", fmt="$#,##0"),
       dict(h="Customers paid", w=11, kind="in", fmt="#,##0"), dict(h="Amount paid ($)", w=14, kind="in", fmt="$#,##0"),
       dict(h="% customers paid", w=10, kind="calc", fmt="0%", formula='IF(OR({F}{r}="",{F}{r}=0,{I}{r}=""),"",{I}{r}/{F}{r})'),
       dict(h="% of compensation paid", w=10, kind="calc", fmt="0%", formula='IF(OR({H}{r}="",{H}{r}=0,{J}{r}=""),"",{J}{r}/{H}{r})'),
       dict(h="Customers not yet paid", w=11, kind="calc", fmt="#,##0", formula='IF({F}{r}="","",{F}{r}-N({I}{r}))'),
       dict(h="Residual approach", w=22, kind="list", lst="Residual approach"), dict(h="Status", w=14, kind="list", lst="Status"),
       dict(h="Target completion", w=12, kind="in", fmt="DD/MM/YYYY"),
       dict(h="Days open", w=9, kind="calc", formula='IF({D}{r}="","",IF({O}{r}="Closed",IF({T}{r}="","",{T}{r}-{D}{r}),TODAY()-{D}{r}))'),
       dict(h="Overdue?", w=10, kind="calc", formula='IF(OR({P}{r}="",{O}{r}="Closed"),"",IF({P}{r}<TODAY(),"OVERDUE",IF({P}{r}-TODAY()<=30,"Due soon","On track")))'),
       dict(h="Regulator notified?", w=11, kind="list", lst="Yes / No"), dict(h="Date closed", w=12, kind="in", fmt="DD/MM/YYYY"),
       dict(h="Assurance over calculations and payments?", w=13, kind="list", lst="Yes / No"), dict(h="Root cause fixed?", w=11, kind="list", lst="Yes / No"),
       dict(h="Program owner", w=16, kind="in")],
 examples=[{"Program ID": "REM-01", "Linked breach or incident": "BR-2026-014", "Issue description": "Administration fee rate 0.05% higher than PDS rate for 8 months after a system change.", "Date identified": D(2026, 5, 12), "Period affected": "Sep 2025 to Apr 2026", "Customers affected": 12000, "Method and beneficial assumptions": "Recalculate fees at PDS rate for each member; add fund earnings on the overcharge; where transaction data is incomplete, assume the member was charged for the full period.", "Estimated compensation incl. interest ($)": 1260000, "Customers paid": 9800, "Amount paid ($)": 1010000, "Residual approach": "Not yet decided", "Status": "Paying", "Target completion": D(2026, 11, 30), "Regulator notified?": "Yes", "Assurance over calculations and payments?": "In progress", "Root cause fixed?": "Yes", "Program owner": "Head of Member Services"}],
 dashboard=[("Programs recorded", lambda R: f"=COUNTA({R('Program ID')})"),
            ("Open programs", lambda R: f"=COUNTA({R('Program ID')})-COUNTIF({R('Status')},\"Closed\")"),
            ("Programs overdue", lambda R: f"=COUNTIF({R('Overdue?')},\"OVERDUE\")"),
            ("Customers affected", lambda R: f"=SUM({R('Customers affected')})"),
            ("Customers paid", lambda R: f"=SUM({R('Customers paid')})"),
            ("Estimated compensation ($)", lambda R: f"=SUM({R('Estimated compensation incl. interest ($)')})"),
            ("Amount paid ($)", lambda R: f"=SUM({R('Amount paid ($)')})"),
            ("Root cause not yet fixed", lambda R: f"=COUNTIF({R('Root cause fixed?')},\"No\")+COUNTIF({R('Root cause fixed?')},\"In progress\")")],
 breakdowns=[("Programs by status", "Status", "Status")],
 extra=lambda wb, h: rag(wb["Remediation tracker"], f"{h['letters']['Overdue?']}{h['first']}:{h['letters']['Overdue?']}{h['last']}")))
